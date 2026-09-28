"""Draw a deterministic engineering diagram and embed it in the intake document."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

root=Path(__file__).resolve().parents[1]
assets=root/'artifacts/router_specification_v04_qa'
assets.mkdir(parents=True,exist_ok=True)
im=Image.new('RGB',(1950,720),'white');d=ImageDraw.Draw(im)
regular='C:/Windows/Fonts/malgun.ttf';bold='C:/Windows/Fonts/malgunbd.ttf'
def text(x,y,t,size=30,color='#26364B',strong=False):
    font=ImageFont.truetype(bold if strong else regular,size)
    d.text((x,y),t,font=font,fill=color,anchor='mm')
blue='#296399';light='#EDF4FA';red='#B44444'
# External networks have no cable to the local router.
d.rounded_rectangle((660,12,1290,112),radius=18,fill='#FBEEEE',outline='#DBBABA',width=3)
text(975,45,'병원망 · 인터넷',34,red,True)
text(975,85,'연결 없음 / WAN 포트 미사용',27,red)
d.line((975,120,975,225),fill='#D5A4A4',width=4)
d.line((955,153,995,193),fill=red,width=7);d.line((995,153,955,193),fill=red,width=7)
for x in [30,725,1420]:
    d.rounded_rectangle((x,245,x+500,650),radius=20,fill=light,outline='#A9BED0',width=3)
# PC icon.
d.rounded_rectangle((205,273,355,363),radius=8,fill='white',outline=blue,width=6)
d.line((280,366,280,387),fill=blue,width=6);d.line((235,389,325,389),fill=blue,width=6)
text(280,431,'병원 사용자 단말',36,strong=True)
text(280,479,'기존 PC · 반입 대상 제외',29)
text(280,527,'익명화된 슬라이드 업로드',29)
text(280,574,'브라우저로 영상·결과 확인',29)
text(280,620,'IP 자동 할당(DHCP)',27,blue)
# Router icon.
d.rounded_rectangle((890,317,1060,377),radius=12,fill='white',outline=blue,width=6)
for x in [910,946,982,1018]:d.rectangle((x,337,x+19,355),fill=blue)
text(975,431,'독립 공유기',36,strong=True)
text(975,479,'LAN 포트 유선 연결',29)
text(975,527,'DHCP · 서버 주소 예약',29)
text(975,574,'WAN 미연결',29,red,True)
text(975,620,'Wi-Fi 비활성화',27,blue)
# Server icon.
d.rounded_rectangle((1610,272,1730,389),radius=8,fill='white',outline=blue,width=6)
for y in [288,320,352]:
    d.rectangle((1625,y,1714,y+19),outline=blue,width=3)
    d.ellipse((1700,y+5,1708,y+13),fill=blue)
text(1670,431,'반입 로컬 AI 서버',36,strong=True)
text(1670,479,'MeDIAuto 영상 뷰어',29)
text(1670,527,'PD-L1 분석 · 결과 저장',29)
text(1670,574,'원본·환자 대응표 업로드 제외',26)
text(1670,620,'DHCP 예약 주소 사용',27,blue)
for a,b in [(540,715),(1235,1410)]:
    d.line((a,390,b,390),fill=blue,width=6)
    d.polygon([(a,390),(a+19,378),(a+19,402)],fill=blue)
    d.polygon([(b,390),(b-19,378),(b-19,402)],fill=blue)
    text((a+b)//2,352,'LAN',29,blue,True)
    text((a+b)//2,429,'유선',27)
text(975,690,'실선 양방향 화살표: 로컬 데이터 통신    /    × 표시: 외부망 미연결',28)
png=assets/'network_diagram.png';im.save(png)
doc=Document(root/'output/docx/MeDIAuto_PDL1_Server_Router_Specification_v03.docx')
diagram=next(t for t in doc.tables if t.rows[0].cells[0].text=='병원 사용자 단말')
el=OxmlElement('w:p');diagram._tbl.addprevious(el)
from docx.text.paragraph import Paragraph
p=Paragraph(el,doc._body);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4)
pic=p.add_run().add_picture(str(png),width=Inches(6.5))
pic._inline.docPr.set('descr','병원 기존 PC와 AI 서버가 독립 공유기의 LAN 포트로 연결됨. 공유기의 WAN은 병원망·인터넷과 연결하지 않음.')
diagram._tbl.getparent().remove(diagram._tbl)
for p in doc.paragraphs:
    for r in p.runs:
        if 'v0.3' in r.text:r.text=r.text.replace('v0.3','v0.4')
out=root/'output/docx/MeDIAuto_PDL1_Server_Router_Diagram_v04.docx'
doc.save(out);print(out)
