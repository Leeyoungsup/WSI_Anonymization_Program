"""Revise v0.6 for the agreed professor-PC/direct local-server deployment."""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
src = ROOT/'output/docx/MeDIAuto_AI_Project_Privacy_Checklist_Integrated_v06.docx'
out = ROOT/'output/docx/MeDIAuto_AI_Project_Privacy_Checklist_Direct_Local_v07.docx'
doc = Document(src)

def set_text(p, value):
    if p.runs:
        p.runs[0].text = value
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(value)

changes = {
'01-2': '교수님 PC의 승인된 계정으로 익명화 프로그램을 실행한다. 원본·처리본·CSV 폴더가 OneDrive 등 자동 동기화 위치가 아닌지 확인하고 접근권한을 제한한다.',
'01-4': '원본과 환자 결과를 연결하는 대응표는 병원이 지정한 별도 보호 위치에 둔다. AI 서버에는 업로드하지 않는다. 임의 출력 파일명 생성은 환자별 연구번호 관리 기능이 아니다.',
'06-1': '교수님 PC에서 익명화한 슬라이드를 직접 연결된 로컬 AI 서버에 업로드하는 구성으로 설치한다. 서버는 병원망에 연결하지 않는다. PC·서버의 설치 위치, 관리 주체, 직접 연결 인터페이스와 IP를 기록한다.',
'06-2': '서버 방화벽은 교수님 PC와 필요한 서비스 포트만 허용한다. 승인된 접속 주소·인증·전송 보호, 서버·DB 접근권한, 저장장치·백업 보호와 복구를 확인한다. 직접 연결만으로 접근 통제가 완료된 것으로 보지 않는다.',
'06-3': 'PC의 병원망·Wi-Fi·VPN 등 다른 연결과 브리지·인터넷 연결 공유·라우팅을 점검한다. 서버가 PC를 경유해 병원망에 연결되지 않도록 한다. 인터넷 연결 여부는 별도로 확인하고 기록한다.',
'06-4': '코드의 IP 위치 조회·PDF 라이브러리 외부 로딩을 점검한다. 외부 연결 없이 운영할 경우 필요한 자원을 로컬 제공하고 업로드·분석·PDF를 시험한다. 비승인 계정의 조회와 원격 접속을 제한한다.',
'07-3': '참여 종료·업무 변경 시 계정을 비활성화하거나 역할을 회수한다. 어반의 유지보수는 승인된 현장 작업을 기본으로 하며, 원격 접속이 필요하면 별도 승인·연결 검토·기간 제한·기록 후 수행한다.',
'08-2': '원본→익명화 결과→등록용 파일명의 대응은 병원 지정 보호 위치에서 관리하고 AI 서버에 올리지 않는다. 필요하면 교수님 PC의 검토된 사본을 가명 케이스 코드로 이름 변경하고 CSV 대응을 확인한다.',
'08-5': '교수님 PC에서 직접 연결된 로컬 서버 주소로 접속한다. 승인된 Project·Folder를 선택하고 익명화 검토가 끝난 파일만 Upload한다. 원본 폴더·대응표는 선택하지 않는다. 업로드 자체는 익명화 기능이 아니다.',
'12-2': '교수님 PC에서 Visualize → PDF Export로 보고서를 저장한다. 다운로드 폴더가 승인된 비동기화 위치인지 확인하고, PDF 제목·모델명·점수·대상 파일명을 검토한 결과와 대조한다.',
'13-6': '추가 학습이 필요하면 서버·담당자·반출 범위를 별도 승인한다. 현재 직접 연결 구성에 외부 학습 서버·클라우드 이용은 포함하지 않는다. 산출 모델·로그·중간 파일의 식별 위험과 보관·파기 계획을 검토한다.',
'15-2': '교수님 PC의 원본·처리본·CSV·PDF, 별도 대응표, 서버의 영상·임상정보·AI 캐시·주석/패치·감사기록을 구분하여 승인된 보존기간과 파기일을 기록한다.',
'15-3': '종료 시 승인된 범위의 PC 연구 사본·다운로드와 서버 영상·DB·캐시·주석/패치를 파기하고 백업 처리도 확인한다. 병원 원본과 대응표는 별도 보존 기준을 확인해 처리한다.',
}
found=set()
for table in doc.tables:
    for row in table.rows:
        code=row.cells[0].text.strip()
        if code in changes:
            set_text(row.cells[1].paragraphs[0], changes[code])
            found.add(code)
        if code=='06-1':
            p=row.cells[2].paragraphs[0]
            for r in p.runs:
                r.text=r.text.replace('배포 위치·운영/접근 주체','설치 위치·담당자·IP/포트')
        if code=='06-3':
            for r in row.cells[2].paragraphs[0].runs:
                r.text=r.text.replace('외부 통신 허용/차단 기록','망 구성·우회 경로 점검 기록')
        if code=='06-4':
            for r in row.cells[2].paragraphs[0].runs:
                r.text=r.text.replace('시험 계정·접근 범위·조치','오프라인 기능·접근 시험')

assert found==set(changes), set(changes)-found
for p in doc.paragraphs:
    t=p.text
    if t.startswith('적용 범위:'):
        set_text(p, '운영 구성: 교수님 PC에서 WSI 익명화·검토 → 직접 연결한 로컬 서버에 업로드 → MeDIAuto AI 분석·주석 → 교수님 PC에서 결과 확인. AI 서버는 병원망에 연결하지 않습니다. 익명화 도구 1.10.2·Studio 3.3.2 기준이며 미사용 기능은 해당 없음으로 기록합니다.')
    elif t=='06. AI 운영환경 승인':
        set_text(p,'06. PC 직접 연결·로컬 서버 점검')
    elif t.startswith('담당: 병원 전산·개인정보 담당 / 시점:'):
        set_text(p,'구성: 교수님 PC(익명화) → 직접 연결 → 로컬 AI 서버(병원망 미연결)\n담당: 설치 관리자·병원 담당 / 최초 업로드 전 및 연결 변경 시 확인')
    elif t.startswith('조건부: 병원 밖으로 자료를 전달할 때'):
        set_text(p,'기본 운영에는 외부 전달이 없습니다. 교수님 PC → 로컬 서버 업로드는 08단계입니다.\n이 페이지는 어반 등 외부 기관에 별도 제공하는 경우에만 적용하며, 미실시 시 해당 없음으로 기록합니다.')
    elif t.startswith('확인 기준: Studio version.json'):
        set_text(p,'운영 계획은 사용자 설명을 반영했습니다. 병원망 미연결과 인터넷 미연결은 별개이며 실제 연결 상태를 점검합니다. 코드 기준: Studio 3.3.2 / 8e3d0e849e4c. 설치·망 분리·기능 시험 완료를 의미하지 않습니다.')

for section in doc.sections:
    for p in list(section.header.paragraphs)+list(section.footer.paragraphs):
        for r in p.runs:
            if 'v0.6' in r.text:
                r.text=r.text.replace('v0.6','v0.7')
for p in doc.paragraphs:
    for r in p.runs:
        if 'v0.6' in r.text:
            r.text=r.text.replace('v0.6','v0.7')
doc.core_properties.title='MeDIAuto AI 프로젝트 개인정보보호·가명처리 체크리스트 — PC 직접 연결 로컬 서버'
doc.save(out)
print(out)
