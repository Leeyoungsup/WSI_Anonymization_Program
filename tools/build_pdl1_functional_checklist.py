"""Apply institutional, function-oriented terminology without changing scope."""
from pathlib import Path
from docx import Document

root=Path(__file__).resolve().parents[1]
src=root/'output/docx/MeDIAuto_PDL1_Privacy_Checklist_Linked_v09.docx'
out=root/'output/docx/MeDIAuto_PDL1_Functional_Checklist_v10.docx'
doc=Document(src)
changes={
'PD-L1 연구\n개인정보보호 체크리스트':'MeDIAuto PD-L1\n기능별 개인정보보호 체크리스트',
'교수님 PC 익명화 · 로컬 서버 뷰어 · PD-L1 AI 결과':'슬라이드 익명화 · 영상 열람 · PD-L1 AI 분석 결과 관리',
'운영 구성':'적용 범위 및 운영 환경',
'교수님 PC: 익명화·검토 → 직접 연결로 처리본 업로드\n로컬 서버(병원망 미연결): 뷰어·PD-L1 분석 → 교수님 PC에서 결과 확인':'처리 흐름: 익명화·검증 → 데이터 등록 → 영상 열람·PD-L1 결과 검토\n운영 환경: 익명화 작업용 PC와 로컬 AI 서버 직접 연결, 서버의 병원망 미연결',
'01  시작 전 승인과 연결 확인':'01  처리 환경 및 접근 통제',
'02  교수님 PC에서 익명화·검토':'02  슬라이드 익명화 및 검증',
'03  로컬 서버 업로드와 뷰어 확인':'03  데이터 등록 및 영상 열람',
'04  PD-L1 AI 결과 확인':'04  PD-L1 AI 분석 결과 검토',
'05  보관·문제 대응과 작업 기록':'05  데이터 보관 및 사후 관리',
'06  공식 참고자료와 적용 범위':'06  관련 기준 및 공식 참고자료',
'담당: 교수님 또는 승인된 사용자':'담당: 데이터 등록·열람 권한 보유자',
'담당: 교수님·지정 연구 검토자':'담당: 연구책임자·지정 검토자',
'교수님 PC에서 승인된 로컬 서버 주소로 접속한다.':'승인된 사용자 단말에서 로컬 서버 주소로 접속한다.',
'PC·서버 직접 연결':'연결 구성 관리',
'접근과 저장 보호':'접근·저장 보호',
'설치 기능 확인':'기능 범위 설정',
'대상과 영상 검사':'대상 파일 검사',
'형식과 검증':'형식·무결성 검증',
'처리본만 업로드':'처리본 등록',
'뷰어 열람 확인':'영상 열람 검증',
'결과 대응과 보관':'결과 연계·보관',
'보관 위치와 기간':'보관 위치·기간',
'외부 제공과 지원':'외부 제공·지원',
'종료와 삭제':'이용 종료·파기',
'PC에서의 처리·검토':'익명화 작업용 PC에서의 처리·검토',
'올리지 않는다.':'등록하지 않는다.',
'자동동기화를 끈다.':'자동동기화를 해제한다.',
'망 우회 연결을 막는다.':'망 우회 연결을 차단한다.',
'영상 저장·파일명 변경은 켜고 CSV 원본 파일명 포함은 끈다.':'영상 저장·파일명 변경을 활성화하고 CSV 원본 파일명 포함을 해제한다.',
'배치별 확인 기록':'처리 단위별 확인 기록',
'v0.9':'v1.0',
}
def update(p):
    if p.text in changes:
        value=changes[p.text]
        p.runs[0].text=value
        for r in p.runs[1:]:r.text=''
        return
    for r in p.runs:
        new=r.text
        for a,b in changes.items():new=new.replace(a,b)
        if new!=r.text:r.text=new
for p in doc.paragraphs:update(p)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:update(p)
for s in doc.sections:
    for p in list(s.header.paragraphs)+list(s.footer.paragraphs):update(p)
doc.core_properties.title='MeDIAuto PD-L1 기능별 개인정보보호 체크리스트'
doc.core_properties.subject='슬라이드 익명화·검증, 데이터 등록·영상 열람, PD-L1 AI 분석 결과 관리'
doc.save(out)
assert len(doc.inline_shapes)==3
assert not any('교수님' in p.text for p in doc.paragraphs)
print(out)
