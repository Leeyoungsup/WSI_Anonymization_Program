"""Approved standalone router topology, without choosing unconfirmed hardware."""
from pathlib import Path
from docx import Document

root=Path(__file__).resolve().parents[1]
doc=Document(root/'output/docx/MeDIAuto_PDL1_Server_Specification_Network_v02.docx')
out=root/'output/docx/MeDIAuto_PDL1_Server_Router_Specification_v03.docx'
replacements={
'v0.2':'v0.3',
'MeDIAuto AI 로컬 서버 및 연결 부속품. 병원 기존 PC는 반입 장비와 사양 기재 대상에서 제외':'MeDIAuto AI 로컬 서버·독립 공유기·연결 케이블. 병원 기존 PC는 반입 장비 및 사양 기재 대상에서 제외',
'병원 사용자 단말과 반입 로컬 AI 서버 직접 연결. AI 서버는 병원망에 연결하지 않음':'독립 공유기의 LAN 포트에 병원 사용자 단말과 AI 서버를 유선 연결. WAN 미연결, 병원망·인터넷 미사용',
'연결 부속품':'공유기 / 연결 부속품',
'선정 후 확정 — 케이블·어댑터 등 구성품':'공유기 1대 사용. 제조사·모델·LAN 규격·케이블 사양은 선정 후 확정. DHCP·주소 예약·Wi-Fi 해제 지원',
'병원 사용자 단말은 연결 상대를 표시한 것이며 반입 장비·사양 기재 대상이 아님':'확정 구성: 독립 공유기 기반 유선 로컬망 / 병원 사용자 단말은 반입 장비에서 제외',
'직접 연결':'독립 공유기',
'처리본 업로드 →\n← 영상·AI 결과\n\n연결 규격: 선정 후 확정':'LAN ↔ LAN\n처리본 업로드 →\n← 영상·AI 결과\n\nWAN 미연결',
'원본·환자 대응표는\n서버 업로드 대상 제외':'DHCP 자동 할당\n서버 주소 예약\nWi-Fi 비활성화',
'병원망 미연결\n외부 통신 여부 별도 확인':'병원망·인터넷 미연결\n원본·환자 대응표 제외',
'병원망  ──  AI 서버 연결 없음':'병원망·인터넷: 연결 없음 / 공유기 WAN 포트: 미사용',
'사용자 단말 → 직접 연결된 로컬 서버 → 영상 열람·PD-L1 결과 확인':'병원 사용자 단말 ↔ 공유기 LAN ↔ AI 서버. 브라우저에서 영상·PD-L1 결과 확인',
'병원 사용자 단말과 AI 서버 직접 연결. 연결 매체·규격·속도는 서버 및 부속품 선정 후 확정':'공유기 LAN 포트에 사용자 단말과 AI 서버를 각각 유선 연결. Wi-Fi는 비활성화. 포트 속도·케이블 규격은 선정 후 확정',
'단말·서버 IP 및 서브넷: 설치 시 확정. 기존 사용 주소와의 충돌 여부 확인':'공유기 DHCP로 단말 주소 자동 할당. 서버는 MAC 주소 기반 DHCP 예약 사용. 실제 LAN 대역·서버 예약 주소는 충돌 확인 후 설치 시 기록',
'AI 서버는 병원망 미연결. 사용자 단말을 통한 우회 접속 경로 점검':'공유기 WAN 및 LAN을 병원망에 연결하지 않음. 사용자 단말의 브리지·연결 공유·라우팅을 통한 우회 연결 차단',
'사용 여부 및 설정값: 설치 시 확정. 병원망 미연결을 인터넷 미연결로 간주하지 않음':'외부 인터넷·외부 DNS 미사용. WAN 설정 불필요. DHCP가 공유기 주소를 게이트웨이/DNS로 제공해도 외부 연결은 구성하지 않음',
'승인된 사용자 단말과 필요한 서비스만 허용하도록 설정 계획. 실제 규칙은 설치 시 확정':'서버 방화벽에서 승인 단말·필요 서비스만 허용. 공유기 관리 암호 설정, 원격 관리·UPnP 비활성화. 실제 서비스 포트는 배포 시 기록',
'현장 지원을 기본 운영안으로 검토. 원격 접속이 필요한 경우 방법·기간·범위와 승인 절차를 별도 확정':'현장 유지보수로 운영하며 상시 원격 접속은 미사용. 업데이트 자료는 승인된 절차로 반입하고 적용 이력 기록',
'인터넷 없이 운영하는 경우 필요한 라이브러리·모델·자원의 로컬 제공과 기능 동작 확인':'인터넷 미연결 상태에서 동작하도록 라이브러리·모델·필요 자원을 로컬에 준비. 외부 통신 의존 제거 및 기능 시험',
'직접 연결·IP/포트·방화벽 및 병원망 우회 경로 점검':'WAN 미연결·LAN 통신·DHCP/서버 예약·Wi-Fi 해제·방화벽 점검',
'제조사 사양표·장비 구성품·외형·전원·설치 공간 확인':'서버·공유기 사양표·케이블·외형·전원·설치 공간 확인',
}
# Match long strings before generic terminology so replacements remain atomic.
ordered=sorted(replacements.items(),key=lambda pair:len(pair[0]),reverse=True)
def change(p):
    for r in p.runs:
        t=r.text
        for old,new in ordered:t=t.replace(old,new)
        if t!=r.text:r.text=t
for p in doc.paragraphs:change(p)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:change(p)
for s in doc.sections:
    for p in list(s.header.paragraphs)+list(s.footer.paragraphs):change(p)
doc.core_properties.subject='반입 AI 서버·독립 공유기 / WAN 미연결 유선 로컬망'
doc.save(out)
print(out)
